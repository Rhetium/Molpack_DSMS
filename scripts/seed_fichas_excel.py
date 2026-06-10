"""
Script de importación de fichas técnicas desde Excel.

Lee la hoja Fichas_Tecnicas del archivo de plantilla, busca el material
por nombre_corporativo, crea cada ficha vía FichaService (con auditoría y
embedding completos) y luego actualiza el estado a Vigente directamente en
la DB (para evitar la restricción único-vigente-por-material+pais, que es
una regla operacional, no una restricción de la DB).

Uso:
    cd <raíz del proyecto>
    python scripts/seed_fichas_excel.py "ruta/al/archivo.xlsx"
"""

import asyncio
import sys
import os
import uuid
from datetime import datetime

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from dotenv import load_dotenv
load_dotenv()

import openpyxl
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from sqlalchemy import select, update, and_

from app.models.kitem import KItem
from app.models.ficha import FichaTecnica
from app.models.material import MaterialComercial
from app.models.kitem_auditoria import KItemAuditoria
from app.models.anomalia import AnomaliaRegistro
from app.schemas.ficha import (
    FichaTecnicaCreateSchema,
    CaracteristicasSchema,
    CaracteristicasContenidoSchema,
    EmpaqueEstibaSchema,
    ManejoDisposicionSchema,
)
from app.services.fichas_services import FichaService
from app.core.dsms_constants import (
    ESTADO_VIGENTE,
    ESTADO_BORRADOR,
    KTYPE_FICHA_TECNICA,
    ESTADO_ABREVIATURAS,
)

USUARIO_IMPORTACION = "importacion_excel"

# Mapeo manual para nombres que no coinciden exactamente
NOMBRE_CORPORATIVO_OVERRIDE = {
    "PORTAVASOS BI PARTIDO PV2  2 CAVIDADES (F) - (BOLSA 4 RUMAS)": "Porta Vasos 2 Cav",
}


# ---------------------------------------------------------------------------
# Helpers de conversión
# ---------------------------------------------------------------------------

def _num(val):
    """Convierte N/C, N/A, None a None; otherwise float."""
    if val is None:
        return None
    s = str(val).strip().upper()
    if s in ("N/C", "N/A", ""):
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _str(val):
    """Convierte N/C, N/A, None a None; otherwise stripped string."""
    if val is None:
        return None
    s = str(val).strip()
    if s.upper() in ("N/C", "N/A", ""):
        return None
    return s


def _int(val):
    """Convierte a int o None."""
    n = _num(val)
    return int(n) if n is not None else None


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------

def leer_fichas_excel(ruta: str) -> list[dict]:
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["Fichas_Tecnicas"]
    headers = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(4, ws.max_row + 1):
        row = {}
        for c in range(1, ws.max_column + 1):
            h = headers[c - 1]
            if h is not None:
                row[h] = ws.cell(r, c).value
        if row.get("nombre_corporativo"):
            rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# Mapeo de fila Excel → esquemas Pydantic
# ---------------------------------------------------------------------------

def construir_caracteristicas(row: dict) -> CaracteristicasSchema | None:
    data = {}

    largo_v = _num(row.get("largo_valor"))
    if largo_v is not None:
        data["dimensiones_largo_valor"] = largo_v
        data["dimensiones_largo_tolerancia"] = _num(row.get("largo_tolerancia"))
        data["dimensiones_largo_unidad"] = _str(row.get("largo_unidad"))

    ancho_v = _num(row.get("ancho_valor"))
    if ancho_v is not None:
        data["dimensiones_ancho_valor"] = ancho_v
        data["dimensiones_ancho_tolerancia"] = _num(row.get("ancho_tolerancia"))
        data["dimensiones_ancho_unidad"] = _str(row.get("ancho_unidad"))

    alto_v = _num(row.get("alto_valor"))
    if alto_v is not None:
        data["dimensiones_alto_valor"] = alto_v
        data["dimensiones_alto_tolerancia"] = _num(row.get("alto_tolerancia"))
        data["dimensiones_alto_unidad"] = _str(row.get("alto_unidad"))

    peso_v = _num(row.get("peso_valor"))
    if peso_v is not None:
        data["peso_valor"] = peso_v
        data["peso_tolerancia"] = _num(row.get("peso_tolerancia"))
        data["peso_unidad"] = _str(row.get("peso_unidad"))

    ruptura_v = _num(row.get("ruptura_valor"))
    if ruptura_v is not None:
        data["ruptura_valor"] = ruptura_v
        data["ruptura_unidad"] = _str(row.get("ruptura_unidad"))

    tc_v = _num(row.get("tiempo_encolado_valor"))
    if tc_v is not None:
        data["tiempo_encolado_valor"] = tc_v
        data["tiempo_encolado_unidad"] = _str(row.get("tiempo_encolado_unidad"))

    abs_v = _num(row.get("absorcion_valor"))
    if abs_v is not None:
        data["porcentaje_absorcion_valor"] = abs_v
        data["porcentaje_absorcion_tolerancia"] = _num(row.get("absorcion_tolerancia"))
        data["porcentaje_absorcion_unidad"] = _str(row.get("absorcion_unidad"))

    def_int_v = _num(row.get("deflexion_int_valor"))
    if def_int_v is not None:
        data["deflexion_interna_valor"] = def_int_v
        data["deflexion_interna_unidad"] = _str(row.get("deflexion_int_unidad"))

    def_ext_v = _num(row.get("deflexion_ext_valor"))
    if def_ext_v is not None:
        data["deflexion_externa_valor"] = def_ext_v
        data["deflexion_externa_unidad"] = _str(row.get("deflexion_ext_unidad"))

    color = _str(row.get("color"))
    if color:
        data["color"] = color

    if not data:
        return None
    return CaracteristicasSchema(**data)


def construir_caracteristicas_contenido(row: dict) -> CaracteristicasContenidoSchema | None:
    data = {}

    pp_v = _num(row.get("profundidad_pilar_valor"))
    if pp_v is not None:
        data["profundidad_pilar_valor"] = pp_v
        data["profundidad_pilar_tolerancia"] = _num(row.get("profundidad_pilar_tolerancia"))
        data["profundidad_pilar_unidad"] = _str(row.get("profundidad_pilar_unidad"))

    da_v = _num(row.get("diametro_alveolo_valor"))
    if da_v is not None:
        data["diametro_alveolo_valor"] = da_v
        data["diametro_alveolo_tolerancia"] = _num(row.get("diametro_alveolo_tolerancia"))
        data["diametro_alveolo_unidad"] = _str(row.get("diametro_alveolo_unidad"))

    pc_v = _num(row.get("profundidad_cavidad_valor"))
    if pc_v is not None:
        data["profundidad_cavidad_valor"] = pc_v
        data["profundidad_cavidad_tolerancia"] = _num(row.get("profundidad_cavidad_tolerancia"))
        data["profundidad_cavidad_unidad"] = _str(row.get("profundidad_cavidad_unidad"))

    dc_v = _num(row.get("diametro_cavidad_valor"))
    if dc_v is not None:
        data["diametro_cavidad_valor"] = dc_v
        data["diametro_cavidad_tolerancia"] = _num(row.get("diametro_cavidad_tolerancia"))
        data["diametro_cavidad_unidad"] = _str(row.get("diametro_cavidad_unidad"))

    if not data:
        return None
    return CaracteristicasContenidoSchema(**data)


def construir_empaque_estiba(row: dict) -> EmpaqueEstibaSchema | None:
    tipo = _str(row.get("tipo_empaque"))
    if not tipo:
        return None

    data = {"tipo_empaque": tipo}

    color_emp = _str(row.get("color_empaque"))
    if color_emp:
        data["color_empaque"] = color_emp

    alto_v = _num(row.get("alto_empaque_valor"))
    if alto_v is not None:
        data["alto_empaque_valor"] = alto_v
        data["alto_empaque_tolerancia"] = _num(row.get("alto_empaque_tolerancia"))
        data["alto_empaque_unidad"] = _str(row.get("alto_empaque_unidad"))

    peso_v = _num(row.get("peso_empaque_valor"))
    if peso_v is not None:
        data["peso_empaque_valor"] = peso_v
        data["peso_empaque_tolerancia"] = _num(row.get("peso_empaque_tolerancia"))
        data["peso_empaque_unidad"] = _str(row.get("peso_empaque_unidad"))

    unidades = _int(row.get("unidades_empaque"))
    if unidades is not None:
        data["undidades_empaque"] = unidades  # typo en schema original

    emp_estiba = _int(row.get("empaques_estiba"))
    if emp_estiba is not None:
        data["empaques_estiba"] = emp_estiba

    camas = _int(row.get("camas_estiba"))
    if camas is not None:
        data["camas_estiba"] = camas

    emp_cama = _int(row.get("empaques_cama"))
    if emp_cama is not None:
        data["empaques_camas_estiba"] = emp_cama

    return EmpaqueEstibaSchema(**data)


def construir_manejo_disposicion(row: dict) -> ManejoDisposicionSchema | None:
    uso = _str(row.get("uso (*)"))
    almacenamiento = _str(row.get("almacenamiento (*)"))
    transporte = _str(row.get("transporte (*)"))
    vida_util = _str(row.get("vida_util (*)"))

    if not all([uso, almacenamiento, transporte, vida_util]):
        return None

    manejo = _str(row.get("manejo (*)")) or "Manipular con cuidado. Consulte al fabricante para instrucciones específicas de manejo."

    data = {
        "uso": uso,
        "manejo": manejo,
        "almacenamiento": almacenamiento,
        "transporte": transporte,
        "vida_util": vida_util,
    }

    inocuidad = _str(row.get("inocuidad"))
    if inocuidad:
        data["inocuidad"] = inocuidad

    disposicion = _str(row.get("disposicion"))
    if disposicion:
        data["disposicion_pt"] = disposicion

    garantias = _str(row.get("garantias"))
    if garantias:
        data["garantias"] = garantias

    manipulacion = _str(row.get("manipulacion"))
    if manipulacion:
        data["manipulacion"] = manipulacion

    return ManejoDisposicionSchema(**data)


# ---------------------------------------------------------------------------
# Avanzar estado a Vigente directamente en la DB
# ---------------------------------------------------------------------------

async def avanzar_a_vigente(session: AsyncSession, ficha: FichaTecnica) -> None:
    now = datetime.now()
    codigo_vigente = (
        f"FT-{ficha.codigo_material_local}-{ficha.pais}-VIG-V{ficha.codigo_version}"
    )

    # Resolver anomalías pendientes para no bloquear el cambio de estado
    await session.execute(
        update(AnomaliaRegistro)
        .where(
            and_(
                AnomaliaRegistro.kitem_id == ficha.id_ficha,
                AnomaliaRegistro.estado == "pendiente",
            )
        )
        .values(
            estado="aceptada",
            resuelto_por=USUARIO_IMPORTACION,
            fecha_resolucion=now,
            nota_resolucion="Aceptada automáticamente durante importación desde Excel.",
        )
    )

    # Actualizar FichaTecnica
    ficha.estado_ficha = ESTADO_VIGENTE
    ficha.codigo_ficha_local = codigo_vigente
    ficha.usuario_ultima_actualizacion = USUARIO_IMPORTACION
    ficha.fecha_actualizacion = now
    session.add(ficha)

    # Actualizar KItem
    await session.execute(
        update(KItem)
        .where(KItem.id == ficha.id_ficha)
        .values(
            estado=ESTADO_VIGENTE,
            metadata_extra={
                "codigo_ficha_local": codigo_vigente,
                "pais": ficha.pais,
                "version": ficha.codigo_version,
            },
            usuario_ultima_actualizacion=USUARIO_IMPORTACION,
            fecha_actualizacion=now,
        )
    )

    # Registro de auditoría
    auditoria = KItemAuditoria(
        id=uuid.uuid4(),
        kitem_id=ficha.id_ficha,
        ktype=KTYPE_FICHA_TECNICA,
        accion="CAMBIO_ESTADO",
        estado_anterior=ESTADO_BORRADOR,
        estado_nuevo=ESTADO_VIGENTE,
        usuario=USUARIO_IMPORTACION,
        fecha=now,
        detalles={
            "codigo_ficha_local": codigo_vigente,
            "codigo_material_local": ficha.codigo_material_local,
            "pais": ficha.pais,
            "version": ficha.codigo_version,
            "transicion": f"{ESTADO_BORRADOR} -> {ESTADO_VIGENTE}",
            "nota": "Avance directo Borrador→Vigente durante importación desde Excel.",
        },
    )
    session.add(auditoria)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

async def main(ruta_excel: str):
    from app.core.database import DATABASE_URL

    engine = create_async_engine(DATABASE_URL)
    SessionLocal = sessionmaker(bind=engine, class_=AsyncSession, expire_on_commit=False)

    # Leer Excel
    filas = leer_fichas_excel(ruta_excel)
    print(f"\nFichas a importar: {len(filas)}\n")

    # Cargar materiales de la DB para buscar por nombre
    async with SessionLocal() as session:
        result = await session.execute(
            select(MaterialComercial.id_material_corporativo, MaterialComercial.nombre_corporativo)
        )
        materiales = {nombre: id_ for id_, nombre in result.fetchall()}

    print("Materiales encontrados en DB:")
    for nombre, id_ in materiales.items():
        print(f"  {id_}  ->  {nombre}")
    print()

    resultados = []

    for i, row in enumerate(filas, 1):
        nombre_corp = row.get("nombre_corporativo", "").strip()
        # Aplicar override de nombre si aplica
        nombre_lookup = NOMBRE_CORPORATIVO_OVERRIDE.get(nombre_corp, nombre_corp)

        id_material = materiales.get(nombre_lookup)
        if not id_material:
            print(f"[{i}] SKIP — Material no encontrado: '{nombre_corp}' (buscado como '{nombre_lookup}')")
            resultados.append({"ficha": nombre_corp, "estado": "SKIP", "razon": "material no encontrado"})
            continue

        codigo_local = str(row.get("codigo_material_local (*)") or row.get("codigo_material_local", "")).strip()
        nombre_local = _str(row.get("nombre_materia_local (*)") or row.get("nombre_local_material"))
        pais = _str(row.get("pais (*)") or row.get("pais")) or "Venezuela"

        ficha_data = FichaTecnicaCreateSchema(
            id_material_corporativo=id_material,
            codigo_material_local=codigo_local,
            nombre_local_material=nombre_local,
            usuario_creador=USUARIO_IMPORTACION,
            pais=pais,
            caracteristicas=construir_caracteristicas(row),
            caracteristicas_contenido=construir_caracteristicas_contenido(row),
            empaque_estiba=construir_empaque_estiba(row),
            manejo_disposicion=construir_manejo_disposicion(row),
        )

        try:
            async with SessionLocal() as session:
                service = FichaService(db_session=session)
                ficha = await service.crear(ficha_data)
                print(f"[{i}] CREADA  — {ficha.codigo_ficha_local}  (material: {nombre_lookup})")

                # Avanzar a Vigente directamente
                await avanzar_a_vigente(session, ficha)
                await session.commit()
                print(f"[{i}] VIGENTE — {ficha.id_ficha}")
                resultados.append({"ficha": ficha.codigo_ficha_local, "estado": "OK"})

        except Exception as e:
            print(f"[{i}] ERROR   — {nombre_corp} / {codigo_local}: {e}")
            resultados.append({"ficha": codigo_local, "estado": "ERROR", "razon": str(e)})

    print("\n--- Resumen ---")
    ok = [r for r in resultados if r["estado"] == "OK"]
    skip = [r for r in resultados if r["estado"] == "SKIP"]
    error = [r for r in resultados if r["estado"] == "ERROR"]
    print(f"  OK:    {len(ok)}")
    print(f"  SKIP:  {len(skip)}")
    print(f"  ERROR: {len(error)}")
    if error:
        print("\nErrores:")
        for r in error:
            print(f"  {r['ficha']}: {r['razon']}")
    if skip:
        print("\nOmitidos:")
        for r in skip:
            print(f"  {r['ficha']}: {r['razon']}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Uso: python scripts/seed_fichas_excel.py <ruta_excel>")
        sys.exit(1)
    asyncio.run(main(sys.argv[1]))
